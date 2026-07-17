"""Copy each group's bottom-of-sheet strategy details into its intake file,
placed to the right of the schedule table (starting column K)."""
import openpyxl, datetime, re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter, column_index_from_string

SRC = "/root/.claude/uploads/43640e3c-45fd-5c32-bf93-da0a8af84015/d880d13b-Call_Delivery_Strategy__July_2026.xlsx"
INTAKE = "/home/user/mygamemystory/call-delivery/intake/2026-07"

NAVY="1F2A44"; BLUE="2E5AAC"; LTBLUE="DCE6F7"; GREY="F2F4F8"; WHITE="FFFFFF"; BORDER="B8C0D0"
thin=Side(style="thin",color=BORDER); box=Border(left=thin,right=thin,top=thin,bottom=thin)
def fill(h): return PatternFill("solid",fgColor=h)

src = openpyxl.load_workbook(SRC, data_only=True)

SHEET_TO_FILE = {"FE":"FE.xlsx","BE":"BE.xlsx","PCO":"PCO.xlsx","Auto":"Auto.xlsx",
    "Repo":"Repo.xlsx","MOD":"MOD.xlsx","Branch Central":"Branch_Central.xlsx",
    "Branch Vendor":"Branch_Vendor.xlsx"}

def extract(name):
    """Pull notes, lists table, and blitz calendar from rows 34-95 of a group sheet."""
    ws = src[name]
    notes = []          # (text, is_header)
    lists_rows = []     # (list_no, list_name, list_strat)
    blitz = None        # dict(labels=(l1,l2), rows=[(date, day, m1, m2)], total=int)

    # locate 'List #' header
    hdr = None
    for r in range(34, 60):
        for c in range(2, 20):
            v = ws.cell(row=r, column=c).value
            if isinstance(v,str) and v.strip()=='List #':
                hdr = (r, c); break
        if hdr: break
    numcol = namecol = stratcol = None
    if hdr:
        hr, hc = hdr
        numcol = hc
        for c in range(hc+1, hc+8):
            v = ws.cell(row=hr, column=c).value
            if isinstance(v,str) and v.strip()=='Lists': namecol = c
            if isinstance(v,str) and 'Specific' in str(v): stratcol = c
        for r in range(hr+1, hr+12):
            no = ws.cell(row=r, column=numcol).value
            nm = ws.cell(row=r, column=namecol).value if namecol else None
            stv = ws.cell(row=r, column=stratcol).value if stratcol else None
            if no is None and nm is None: continue
            lists_rows.append((no, nm, stv))

    # notes from cols B-D, rows 34-95 (skip cells that belong to the lists table cols)
    skipcols = {c for c in (numcol,namecol,stratcol) if c}
    for r in range(34, 96):
        for c in (2,3,4):
            if c in skipcols: continue
            v = ws.cell(row=r, column=c).value
            if isinstance(v,str) and v.strip() and v.strip()!='#REF!':
                t = v.rstrip()
                is_hdr = t.strip().lower() in ("general strategy","initiatives","inititives")
                notes.append((t, is_hdr))

    # blitz calendar: dates in col W(23), day col X(24), marks Y(25), Z(26)
    brows = []
    for r in range(39, 70):
        d = ws.cell(row=r, column=23).value
        if isinstance(d, datetime.datetime):
            brows.append((d.date(), ws.cell(row=r,column=24).value,
                          ws.cell(row=r,column=25).value, ws.cell(row=r,column=26).value))
    if brows:
        l1 = ws.cell(row=38, column=25).value or "AM"
        l2 = ws.cell(row=38, column=26).value or "PM"
        total = ws.cell(row=71, column=25).value
        blitz = dict(labels=(str(l1),str(l2)), rows=brows, total=total)
    return notes, lists_rows, blitz

for sheet, fname in SHEET_TO_FILE.items():
    notes, lists_rows, blitz = extract(sheet)
    path = f"{INTAKE}/{fname}"
    wb = openpyxl.load_workbook(path)
    sh = wb["Schedule"]

    C0 = 11  # column K
    r = 2
    # Panel title
    sh.merge_cells(start_row=r, start_column=C0, end_row=r, end_column=C0+3)
    t = sh.cell(row=r, column=C0, value="GROUP STRATEGY & NOTES  (carried over from July 2026 master)")
    t.font=Font(bold=True,size=12,color=WHITE); t.fill=fill(NAVY)
    t.alignment=Alignment(vertical="center",indent=1)
    sh.row_dimensions[r].height=24
    r += 2

    # Notes
    for text, is_hdr in notes:
        c = sh.cell(row=r, column=C0, value=text)
        if is_hdr:
            c.font=Font(bold=True,size=11,color=BLUE)
        else:
            c.font=Font(size=10)
        sh.merge_cells(start_row=r, start_column=C0, end_row=r, end_column=C0+3)
        r += 1
    r += 1

    # Lists table
    if lists_rows:
        c = sh.cell(row=r, column=C0, value="Lists"); c.font=Font(bold=True,size=11,color=BLUE)
        r += 1
        hdrs = ["List #","List","List Specific Strategy"]
        for i,h in enumerate(hdrs):
            c = sh.cell(row=r, column=C0+i, value=h)
            c.font=Font(bold=True,color=WHITE); c.fill=fill(BLUE); c.border=box
            c.alignment=Alignment(horizontal="center",vertical="center")
        r += 1
        for no,nm,stv in lists_rows:
            for i,v in enumerate((no,nm,stv)):
                c = sh.cell(row=r, column=C0+i, value=v); c.border=box
                c.alignment=Alignment(vertical="center", horizontal="center" if i==0 else "left",
                                      indent=0 if i==0 else 1)
                c.font=Font(size=10)
            r += 1
        r += 1

    # Blitz calendar
    if blitz:
        c = sh.cell(row=r, column=C0, value="Blitz Calendar"); c.font=Font(bold=True,size=11,color=BLUE)
        r += 1
        hdrs = ["Date","Day",blitz["labels"][0],blitz["labels"][1]]
        for i,h in enumerate(hdrs):
            c = sh.cell(row=r, column=C0+i, value=h)
            c.font=Font(bold=True,color=WHITE); c.fill=fill(BLUE); c.border=box
            c.alignment=Alignment(horizontal="center",vertical="center")
        r += 1
        for d, day, m1, m2 in blitz["rows"]:
            vals = [d, day, m1, m2]
            for i,v in enumerate(vals):
                c = sh.cell(row=r, column=C0+i, value=v); c.border=box
                c.alignment=Alignment(horizontal="center",vertical="center")
                c.font=Font(size=10)
                if i==0: c.number_format="mm/dd"
                if i>=2 and v: c.fill=fill(LTBLUE); c.font=Font(size=10,bold=True,color=NAVY)
            r += 1
        tr = sh.cell(row=r, column=C0+1, value="Total")
        tr.font=Font(bold=True,color=NAVY); tr.alignment=Alignment(horizontal="right")
        tv = sh.cell(row=r, column=C0+2, value=blitz["total"])
        tv.font=Font(bold=True,color=NAVY); tv.alignment=Alignment(horizontal="center"); tv.border=box
        r += 1

    # column widths for the panel
    for i,w in enumerate([46,16,34,14]):
        sh.column_dimensions[get_column_letter(C0+i)].width=w

    wb.save(path)
    print(f"{sheet}: notes={len(notes)} lists={len(lists_rows)} blitz={'yes ('+str(len(blitz['rows']))+' days)' if blitz else 'no'} -> {fname}")
