"""Add Reference tabs to the master: Voice, AM Voice IDs, Card Hours,
Branch Standing Strategy (1Pay/2Pay), Blitz Rules."""
import openpyxl, datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC="/root/.claude/uploads/43640e3c-45fd-5c32-bf93-da0a8af84015/d880d13b-Call_Delivery_Strategy__July_2026.xlsx"
P="/home/user/mygamemystory/call-delivery/templates/Call_Delivery_MASTER.xlsx"
NAVY="1F2A44"; BLUE="2E5AAC"; LTBLUE="DCE6F7"; GREY="F2F4F8"; WHITE="FFFFFF"; BORDER="B8C0D0"
thin=Side(style="thin",color=BORDER); box=Border(left=thin,right=thin,top=thin,bottom=thin)
def fill(h): return PatternFill("solid",fgColor=h)
F=lambda **kw: Font(name="Arial",**kw)

src=openpyxl.load_workbook(SRC, data_only=True)
wb=openpyxl.load_workbook(P)
for n in ("Ref Voice","Ref AM Voice IDs","Ref Card Hours","Ref Branch Strategy","Ref Blitz Rules","Ref Group Numbers"):
    if n in wb.sheetnames: del wb[n]

# Dialer group numbers per business group (from the original workbook's "Group #'s" column).
# Single source of truth: the Daily Rollup looks these up by Business Group name.
GROUP_NUMBERS = [
    ("FE","14 / 54"), ("BE","12 / 52"), ("PCO","11 / 51"), ("Auto","36 / 76 / 116"),
    ("Repo","17 / 57"), ("MOD","18 / 58"), ("Branch Central","15 / 55"),
    ("Branch Vendor","79 (Vendor)"), ("Card","40 / 80"), ("ARC","7 / 47"),
    ("Optional Products","130 / 131"), ("Call Escalation Team","37 / 77"),
]

def new_sheet(name, title):
    sh=wb.create_sheet(name); sh.sheet_view.showGridLines=False
    sh.sheet_properties.tabColor=LTBLUE
    sh.merge_cells("A1:F1")
    t=sh["A1"]; t.value=title; t.font=F(bold=True,size=12,color=WHITE); t.fill=fill(NAVY)
    t.alignment=Alignment(vertical="center",indent=1); sh.row_dimensions[1].height=22
    return sh

def hdr(sh,row,cols):
    for i,h in enumerate(cols):
        c=sh.cell(row=row,column=1+i,value=h)
        c.font=F(bold=True,color=WHITE,size=10); c.fill=fill(BLUE); c.border=box
        c.alignment=Alignment(horizontal="center")

# ---- Ref Voice ----
sh=new_sheet("Ref Voice","AUTO-MESSAGING VOICE — daily calendar (Daily Rollup reads this)")
hdr(sh,2,["Date","Voice"]) ; sh.cell(row=2,column=1).value="Date"
ws=src["Voice"]; r=2
for i in range(3,34):
    d=ws.cell(row=i,column=1).value; v=ws.cell(row=i,column=2).value
    if not isinstance(d,datetime.datetime): continue
    r+=1
    sh.cell(row=r,column=1,value=d.date()).number_format="mm/dd/yyyy"
    sh.cell(row=r,column=2,value=(v.strip() if isinstance(v,str) else v))
    for c in (1,2): sh.cell(row=r,column=c).border=box; sh.cell(row=r,column=c).font=F(size=10)
sh.column_dimensions["A"].width=13; sh.column_dimensions["B"].width=10

# ---- Ref AM Voice IDs ----
sh=new_sheet("Ref AM Voice IDs","AUTO-MESSAGE CAMPAIGN & VOICE IDs")
ws=src["AMs"]
hdr(sh,2,["Campaign","Lists","Female ID","Male ID"])
r=2
for i in range(2,11):
    vals=[ws.cell(row=i,column=c).value for c in range(1,5)]
    if not vals[0]: continue
    r+=1
    for j,v in enumerate(vals):
        c=sh.cell(row=r,column=1+j,value=v); c.border=box; c.font=F(size=10)
for c,w in {"A":22,"B":22,"C":11,"D":11}.items(): sh.column_dimensions[c].width=w

# ---- Ref Card Hours ----
sh=new_sheet("Ref Card Hours","CARD HOURS — daily open/close (Daily Rollup reads this)")
hdr(sh,2,["Date","Open","Close","Week"])
ws=src["Card Schedule"]; r=2
for i in range(4,35):
    d=ws.cell(row=i,column=2).value
    if not isinstance(d,datetime.datetime): continue
    r+=1
    sh.cell(row=r,column=1,value=d.date()).number_format="mm/dd/yyyy"
    for j,c0 in enumerate((3,4,5)):
        sh.cell(row=r,column=2+j,value=ws.cell(row=i,column=c0).value)
    for c in range(1,5): sh.cell(row=r,column=c).border=box; sh.cell(row=r,column=c).font=F(size=10)
for c,w in {"A":13,"B":10,"C":10,"D":9}.items(): sh.column_dimensions[c].width=w

# ---- Ref Branch Strategy (1Pay + 2Pay stacked) ----
sh=new_sheet("Ref Branch Strategy","BRANCH STANDING STRATEGY — weekly list priorities (1Pay & 2Pay)")
r=2
for tab in ("1Pay Branch Strategy","2Pay Branch Strategy"):
    ws=src[tab]
    r+=1
    c=sh.cell(row=r,column=1,value=tab.upper()); c.font=F(bold=True,size=11,color=BLUE)
    for i in range(1, ws.max_row+1):
        vals=[ws.cell(row=i,column=cc).value for cc in range(1,11)]
        if not any(v is not None and str(v).strip()!="" for v in vals): continue
        r+=1
        for j,v in enumerate(vals):
            if v is None: continue
            c=sh.cell(row=r,column=1+j,value=v); c.font=F(size=9)
            if isinstance(v,str) and v.strip() in ("LIST ID","Priority ","Time Zone","TZ ","Desc"):
                c.font=F(size=9,bold=True,color=WHITE); c.fill=fill(BLUE)
    r+=1
for c,w in {"A":11,"B":24,"C":9,"D":8,"E":40,"F":24,"G":9,"H":10,"I":40}.items():
    sh.column_dimensions[c].width=w

# ---- Ref Blitz Rules ----
sh=new_sheet("Ref Blitz Rules","BLITZ STANDING RULES — windows, recall time, designated lists")
ws=src["Blitz"]; r=2
for i in range(2,19):
    vals=[ws.cell(row=i,column=cc).value for cc in range(2,11)]   # B..J (skip calendar cols L+)
    if not any(v is not None and str(v).strip()!="" for v in vals): continue
    r+=1
    for j,v in enumerate(vals):
        if v is None: continue
        c=sh.cell(row=r,column=1+j,value=v); c.font=F(size=10)
        if isinstance(v,str) and (v in ("FE","BE") or "Evening" in v or "2Pay" in v or "Saturday AM" in v):
            c.font=F(size=10,bold=True,color=NAVY)
sh.cell(row=r+2,column=1,value="Per-date blitz calendars live in each group's intake file (FE/BE notes panel).").font=F(italic=True,size=9,color="6B7280")
for c,w in {"A":6,"B":16,"C":42,"D":6,"E":6,"F":16,"G":42}.items(): sh.column_dimensions[c].width=w

# ---- Ref Group Numbers ----
sh=new_sheet("Ref Group Numbers","DIALER GROUP NUMBERS — per business group (Daily Rollup reads this)")
hdr(sh,2,["Business Group","Group #'s"])
r=2
for name,num in GROUP_NUMBERS:
    r+=1
    c=sh.cell(row=r,column=1,value=name); c.border=box; c.font=F(size=10)
    c=sh.cell(row=r,column=2,value=num); c.border=box; c.font=F(size=10)
    c.alignment=Alignment(horizontal="center")
sh.cell(row=r+2,column=1,value="Edit a group's number here and it updates on the Daily Rollup automatically.").font=F(italic=True,size=9,color="6B7280")
sh.column_dimensions["A"].width=24; sh.column_dimensions["B"].width=16

wb.save(P)
print("reference tabs added:", [n for n in wb.sheetnames])
