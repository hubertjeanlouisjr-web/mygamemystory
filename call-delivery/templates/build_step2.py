"""Step 2: generate pre-filled intake files from July data + master workbook."""
import openpyxl, datetime, re, os
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = "/root/.claude/uploads/43640e3c-45fd-5c32-bf93-da0a8af84015/d880d13b-Call_Delivery_Strategy__July_2026.xlsx"
BASE = "/home/user/mygamemystory/call-delivery"
INTAKE = f"{BASE}/intake/2026-07"
os.makedirs(INTAKE, exist_ok=True)

NAVY="1F2A44"; BLUE="2E5AAC"; LTBLUE="DCE6F7"; GREY="F2F4F8"; WHITE="FFFFFF"; BORDER="B8C0D0"
thin = Side(style="thin", color=BORDER); box = Border(left=thin,right=thin,top=thin,bottom=thin)
def fill(h): return PatternFill("solid", fgColor=h)

# ---------------- parser (validated earlier) ----------------
WAVE = re.compile(r"^\s*([ECMP](?:\s*-\s*[ECMP]){0,3})\s*(?:\((.*?)\))?\s*$")
def parse_time(v):
    if isinstance(v, datetime.time): return v
    if isinstance(v, datetime.datetime): return v.time()
    if not isinstance(v, str): return None
    s = v.strip().lower().replace(' ','')
    m = re.match(r'^(\d{1,2}):?(\d{2})?(a|p|am|pm)?$', s)
    if not m: return None
    h=int(m.group(1)); mm=int(m.group(2) or 0); ap=m.group(3)
    if ap and ap.startswith('p') and h!=12: h+=12
    if ap and ap.startswith('a') and h==12: h=0
    return datetime.time(h,mm) if h<24 else None
def parse_label(lbl):
    if not isinstance(lbl,str) or not lbl.strip(): return None,None
    s=lbl.strip(); m=WAVE.match(s)
    if m:
        wave=m.group(1).replace(' ',''); tag=(m.group(2) or '').strip().rstrip("'s").strip() or None
        if tag and tag.lower().startswith("am"): tag="AMs"
        return wave,tag
    t=s.strip('()').strip()
    if t.upper().startswith('WFO'): return None,'WFO'
    if 'blitz' in t.lower(): return None,'Blitz'
    if 'tsc' in t.lower(): return None,'TSC (All)'
    if 'branch' in t.lower(): return None,'Branch'
    if 'vendor' in t.lower(): return None,'Vendor'
    if t.lower()=='all': return 'All',None
    return None,t

src = openpyxl.load_workbook(SRC, data_only=True)
def parse_sheet(name):
    ws=src[name]; rows=[]
    for r in range(3, ws.max_row+1):
        d=ws.cell(row=r,column=2).value
        if not isinstance(d,datetime.datetime) or d.year!=2026 or d.month!=7: continue
        note=ws.cell(row=r,column=3).value
        note=note.strip() if isinstance(note,str) and note.strip() else None
        prio=0
        for c in range(5, ws.max_column, 2):
            t=parse_time(ws.cell(row=r,column=c).value)
            if t is None: continue
            wave,tag=parse_label(ws.cell(row=r,column=c+1).value or '')
            prio+=1
            rows.append((d.date(), t, wave, tag, prio, note if prio==1 else None))
    return rows

GROUPS = {"FE":"FE","BE":"BE","PCO":"PCO","Auto":"Auto","Repo":"Repo","MOD":"MOD",
          "Branch Central":"Branch Central","Branch Vendor":"Branch Vendor"}
TAGS   = ["AMs","Branch","Blitz","TSC (All)","WFO","Vendor","NCs","1s & Slows","Repo","RP",
          "Recent Pay","Broken Prom","All","Other"]
WAVES  = ["E","C","M","P","E-C","C-M","M-P","E-C-M","C-M-P","E-C-M-P","All"]
ALL_GROUPS = ["OM_All","FE","BE","PCO","MOD","AUTO Direct Collections","Repo","Auto",
    "ARC","CPOD / NCC","CARE (Outbound)","Branch Central","Branch Vendor","Optional Products",
    "West Coast Pilot","Spanish","Card","Call Escalation Team","Sales","Recoveries"]
IDS = ["ADAN","ADAO","ADAR","FEAO","FEAR","BEAO","BEAR","PCAO","PCAR","REAO","REAH",
    "BWAO","BWAR","NCAO","NCAH","LDA3","LDA4","SAMO","SAMR","BR_CAP_CONS_HB","BR_CAP_CONS_LB",
    "BR_CAP_CONS_SLOW_HB","BR_CAP_CONS_SLOW_LB","BR_1PAY_HB","BR_1PAY_LB","BR_SLOW_HB",
    "BR_SLOW_LB","BR_ADHOC_HB","BR_ADHOC_LB","N/A"]

def build_intake(group, rows, path):
    wb = openpyxl.Workbook()
    # Lists
    lists = wb.active; lists.title="Lists"
    for col,(hdr,vals) in {"A":("Business Groups",ALL_GROUPS),"B":("Timezone Waves",WAVES),
                           "C":("Segment / Tag",TAGS),"D":("List ID / Campaign",IDS)}.items():
        c=lists[f"{col}1"]; c.value=hdr; c.font=Font(bold=True,color=WHITE); c.fill=fill(BLUE)
        for i,v in enumerate(vals,start=2): lists[f"{col}{i}"]=v
        lists.column_dimensions[col].width=24
    GR=f"Lists!$A$2:$A${len(ALL_GROUPS)+1}"; WR=f"Lists!$B$2:$B${len(WAVES)+1}"
    TR=f"Lists!$C$2:$C${len(TAGS)+1}"; IR=f"Lists!$D$2:$D${len(IDS)+1}"

    sh = wb.create_sheet("Schedule",0); sh.sheet_view.showGridLines=False
    sh.sheet_properties.tabColor=BLUE
    sh.merge_cells("B2:I2"); t=sh["B2"]
    t.value="CALL DELIVERY  —  Outbound Dialing Schedule (Intake)"
    t.font=Font(bold=True,size=16,color=WHITE); t.fill=fill(NAVY)
    t.alignment=Alignment(vertical="center",indent=1); sh.row_dimensions[2].height=30
    sh.merge_cells("B3:I3"); s=sh["B3"]
    s.value="Pre-filled from the July 2026 master workbook. Review, adjust, and overwrite this file in the Intake folder when your schedule changes.  All times EASTERN."
    s.font=Font(italic=True,size=10,color=NAVY); s.fill=fill(LTBLUE)
    s.alignment=Alignment(vertical="center",indent=1); sh.row_dimensions[3].height=20

    def meta(lc,label,vc,val=None):
        a=sh[lc]; a.value=label; a.font=Font(bold=True,color=NAVY)
        a.alignment=Alignment(horizontal="right",vertical="center")
        b=sh[vc]; b.border=box; b.alignment=Alignment(vertical="center",indent=1)
        if val is not None: b.value=val
    meta("B5","Business Group:","C5",group); sh["C5"].font=Font(bold=True,color=BLUE,size=12)
    meta("E5","Period (Month):","F5","2026-07")
    meta("B6","Submitted By:","C6","(auto-converted)")
    meta("E6","Submitted Date:","F6",datetime.date(2026,7,17))
    sh["F6"].number_format="mm/dd/yyyy"
    for r in (5,6): sh.row_dimensions[r].height=20
    dvg=DataValidation(type="list",formula1=GR,allow_blank=False); sh.add_data_validation(dvg); dvg.add(sh["C5"])

    HDR=9; headers=["Business Group","Date","Start Time (ET)","Timezone Wave",
                    "List ID / Campaign","Priority","Segment / Tag","Special Instructions"]
    for i,h in enumerate(headers):
        c=sh.cell(row=HDR,column=2+i,value=h); c.font=Font(bold=True,color=WHITE)
        c.fill=fill(BLUE); c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
        c.border=box
    sh.row_dimensions[HDR].height=30
    Lc=get_column_letter
    BG,DT,TM,TZ,LI,PR,SG,SI=[Lc(2+i) for i in range(8)]
    n_data=max(len(rows)+50, 120)
    first=HDR+1; last=HDR+n_data
    for idx in range(n_data):
        r=first+idx
        sh[f"{BG}{r}"]=f'=IF($C$5="","",$C$5)'
        sh[f"{BG}{r}"].font=Font(color="6B7280")
        sh[f"{DT}{r}"].number_format="mm/dd/yyyy"; sh[f"{TM}{r}"].number_format="h:mm AM/PM"
        for cc in (BG,DT,TM,TZ,LI,PR,SG,SI):
            sh[f"{cc}{r}"].border=box
            sh[f"{cc}{r}"].alignment=Alignment(vertical="center",
                horizontal="left" if cc==SI else "center", indent=1 if cc==SI else 0)
    for i,(d,t,wave,tag,prio,note) in enumerate(rows):
        r=first+i
        sh[f"{DT}{r}"]=d; sh[f"{TM}{r}"]=t
        if wave: sh[f"{TZ}{r}"]=wave
        sh[f"{LI}{r}"]="N/A"; sh[f"{PR}{r}"]=prio
        if tag: sh[f"{SG}{r}"]=tag
        if note: sh[f"{SI}{r}"]=note
    for c,w in {BG:20,DT:13,TM:14,TZ:14,LI:20,PR:9,SG:15,SI:40,"A":2}.items():
        sh.column_dimensions[c].width=w
    tbl=Table(displayName="tblSchedule",ref=f"{BG}{HDR}:{SI}{last}")
    tbl.tableStyleInfo=TableStyleInfo(name="TableStyleLight9",showRowStripes=True,
        showColumnStripes=False,showFirstColumn=False,showLastColumn=False)
    sh.add_table(tbl)
    def dv(**kw):
        d=DataValidation(**kw); sh.add_data_validation(d); return d
    dv(type="list",formula1=WR,allow_blank=True).add(f"{TZ}{first}:{TZ}{last}")
    dv(type="list",formula1=IR,allow_blank=True).add(f"{LI}{first}:{LI}{last}")
    dv(type="list",formula1=TR,allow_blank=True).add(f"{SG}{first}:{SG}{last}")
    dv(type="whole",operator="between",formula1="1",formula2="20",allow_blank=True).add(f"{PR}{first}:{PR}{last}")
    dv(type="date",operator="greaterThanOrEqual",formula1="DATE(2020,1,1)",allow_blank=True).add(f"{DT}{first}:{DT}{last}")
    dv(type="time",operator="between",formula1="0",formula2="1",allow_blank=True).add(f"{TM}{first}:{TM}{last}")
    sh.freeze_panes=f"{DT}{first}"
    wb.save(path)
    return len(rows)

all_rows = {}
for sheet, group in GROUPS.items():
    rows = parse_sheet(sheet)
    if not rows: continue
    all_rows[group] = rows
    fname = group.replace(" ","_") + ".xlsx"
    n = build_intake(group, rows, f"{INTAKE}/{fname}")
    print(f"{group}: {n} passes -> intake/2026-07/{fname}")
print("TOTAL:", sum(len(v) for v in all_rows.values()))
