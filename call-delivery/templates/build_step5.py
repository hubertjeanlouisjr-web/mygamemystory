"""Build 4 additional intake files (Card, ARC, Optional Products, CET)
and notes panels, reusing build_step2's template builder."""
import sys, os, datetime, importlib.util
S = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location("b2", f"{S}/build_step2.py")
b2 = importlib.util.module_from_spec(spec)
spec.loader.exec_module(b2)   # regenerates the 8 files (idempotent) and gives us build_intake

import openpyxl
src = openpyxl.load_workbook(b2.SRC, data_only=True)
INTAKE = b2.INTAKE

# extend segment tags for the new files
b2.TAGS = b2.TAGS + ["ASAP","First","Second","Third","Second/Third","Third/First",
                     "DSA Expired","DAS Exhausted","DSAR (30+)"]

# ---------- Card (from Card Passes, COLLECTION block rows 2-32) ----------
ws = src["Card Passes"]
rows=[]
for r in range(2, 33):
    d = ws.cell(row=r, column=2).value
    if not isinstance(d, datetime.datetime): continue
    note = None
    prio = 0
    for c in range(3, 41, 2):   # time col C,E,G...; label col +1
        tv = ws.cell(row=r, column=c).value
        lv = ws.cell(row=r, column=c+1).value
        hour = 8 + (c-3)//2
        t = b2.parse_time(tv)
        if t is None and isinstance(tv,str) and tv.strip().lower()=="constant":
            note = "Constant"; continue
        if t is None and isinstance(lv,str) and lv.strip().upper()=="ASAP":
            t = datetime.time(hour,0); tag="ASAP"; wave=None
        elif t is not None:
            wave, tag = b2.parse_label(lv if lv is not None else '')
        else:
            continue
        prio += 1
        rows.append((d.date(), t, wave, tag, prio, note if prio==1 else None))
n=b2.build_intake("Card", rows, f"{INTAKE}/Card.xlsx"); print("Card:", n, "passes")

# ---------- ARC (per-date blocks of 6 sub-strategies) ----------
ws = src["ARC"]
rows=[]; cur=None; prio_by_day={}
for r in range(2, 219):
    a = ws.cell(row=r, column=1).value
    if isinstance(a, datetime.datetime):
        cur = a.date(); continue
    if cur is None or not isinstance(a,str) or not a.strip(): continue
    strat = a.strip()
    for c in range(5, 40, 2):
        t = b2.parse_time(ws.cell(row=r, column=c).value)
        if t is None: continue
        wave, tag = b2.parse_label(ws.cell(row=r, column=c+1).value or '')
        prio_by_day[cur] = prio_by_day.get(cur,0)+1
        rows.append((cur, t, wave, tag, prio_by_day[cur], strat))
n=b2.build_intake("ARC", rows, f"{INTAKE}/ARC.xlsx"); print("ARC:", n, "passes")

# ---------- Optional Products (daily priority calendar) ----------
ws = src["Optional Products"]
PRIOS={"First","Second","Third","Second/Third","Third/First"}
rows=[]
for r in range(3, 34):
    d = ws.cell(row=r, column=2).value
    if not isinstance(d, datetime.datetime): continue
    v = ws.cell(row=r, column=4).value
    if not isinstance(v,str) or not v.strip(): continue
    v=v.strip()
    if v in PRIOS:
        rows.append((d.date(), None, None, v, 1, "Priority Order: "+v))
    else:
        rows.append((d.date(), None, None, None, 1, v))   # e.g. holiday text
n=b2.build_intake("Optional Products", rows, f"{INTAKE}/Optional_Products.xlsx")
print("Optional Products:", n, "rows")

# ---------- CET (schedule empty: old volumes were all broken links) ----------
n=b2.build_intake("Call Escalation Team", [], f"{INTAKE}/CET.xlsx"); print("CET: 0 rows (fresh)")

# ---------- notes panels ----------
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
NAVY="1F2A44"; BLUE="2E5AAC"; WHITE="FFFFFF"; BORDER="B8C0D0"
thin=Side(style="thin",color=BORDER); box=Border(left=thin,right=thin,top=thin,bottom=thin)
def fill(h): return PatternFill("solid",fgColor=h)

def panel(path, lines, lists_rows=None):
    wb=openpyxl.load_workbook(path); sh=wb["Schedule"]; C0=11; r=2
    sh.merge_cells(start_row=r,start_column=C0,end_row=r,end_column=C0+3)
    t=sh.cell(row=r,column=C0,value="GROUP STRATEGY & NOTES  (carried over from July 2026 master)")
    t.font=Font(bold=True,size=12,color=WHITE); t.fill=fill(NAVY)
    t.alignment=Alignment(vertical="center",indent=1); sh.row_dimensions[r].height=24
    r+=2
    for text,is_hdr in lines:
        c=sh.cell(row=r,column=C0,value=text)
        c.font=Font(bold=True,size=11,color=BLUE) if is_hdr else Font(size=10)
        sh.merge_cells(start_row=r,start_column=C0,end_row=r,end_column=C0+3)
        r+=1
    if lists_rows:
        r+=1
        c=sh.cell(row=r,column=C0,value="Lists"); c.font=Font(bold=True,size=11,color=BLUE); r+=1
        for i,h in enumerate(["List #","List","Notes"]):
            c=sh.cell(row=r,column=C0+i,value=h)
            c.font=Font(bold=True,color=WHITE); c.fill=fill(BLUE); c.border=box
            c.alignment=Alignment(horizontal="center")
        r+=1
        for row in lists_rows:
            for i,v in enumerate(row):
                c=sh.cell(row=r,column=C0+i,value=v); c.border=box; c.font=Font(size=10)
            r+=1
    from openpyxl.utils import get_column_letter
    for i,w in enumerate([46,16,34,14]): sh.column_dimensions[get_column_letter(C0+i)].width=w
    wb.save(path)

panel(f"{INTAKE}/Card.xlsx", [
    ("General Strategy",True),
    ("COLLECTION: daily 8:00a full pass + 12:00p pass; afternoon/evening passes run ASAP as capacity allows.",False),
    ("Hours vary by day — see the Ref Card Hours tab in the master workbook (open/close + week number).",False),
    ("PCO and Recoveries card sections had no scheduled passes in July.",False),
    ("Recoveries runs Tues & Thurs when active.",False)])
panel(f"{INTAKE}/ARC.xlsx", [
    ("General Strategy",True),
    ("Each day runs up to six sub-strategies; the sub-strategy for each pass is recorded in",False),
    ("the Special Instructions column: Broken Promise / Recent Pay / Dialer Only, Hit List / APP,",False),
    ("Preplacement / Strat 2, Hit Priority / Strat 7 / APP, Strategy, Agency Recall.",False)])
panel(f"{INTAKE}/Optional_Products.xlsx", [
    ("General Strategy",True),
    ("M-F (1) Full Pass; additional passes as time allows.",False),
    ("Hours of Operation: 8:00am - 6:00pm ET  (Groups 130 / 131).",False),
    ("Daily rows carry the priority order (First / Second / Third) in the Segment / Tag column.",False)],
    lists_rows=[(30000,"Optional Products",""),(30001,"Optional Products",""),(30002,"Optional Products","")])
panel(f"{INTAKE}/CET.xlsx", [
    ("General Strategy",True),
    ("Standing daily priority ladder (volumes were broken links in the old workbook - to be re-supplied):",False),
    ("   Priority 1:  DSA Expired",False),
    ("   Priority 2:  DAS Exhausted",False),
    ("   Priority 3:  DSAR (30+)",False),
    ("Hours of Operation: 8:00a - 5:00p  (Groups 37 / 77).",False),
    ("Schedule left blank - CET to submit passes via this template going forward.",False)])
print("notes panels added")
